"""Excel -> CSV: cada hoja con datos se convierte en un CSV derivado, que luego se analiza como cualquier otro log.

Por qué así y no leyendo el Excel directamente: el resto del pipeline (perfilado, Inspector, ingesta, seudonimización) ya está probado con
CSV; convertir en la entrada mantiene un solo camino. El Excel original queda como evidencia con su hash, y cada CSV derivado registra en la
custodia de qué archivo y de qué hoja sale (ver `projects.Project`).

Reglas de conversión:
* La cabecera es la primera fila con algún valor. Las columnas sin nombre se llaman `col_N`; los nombres repetidos se numeran.
* Las fechas de Excel salen como texto ISO (`2026-10-04 23:02:59`), no como el número de serie interno de Excel.
* Los números enteros guardados como decimales (`80.0`) salen como enteros (`80`); un puerto o un código no debe cambiar de forma.
* Las filas vacías se omiten; una hoja sin filas de datos no genera CSV.
"""
from __future__ import annotations

import csv
from collections.abc import Callable, Iterable, Iterator
from datetime import date, datetime, time
from pathlib import Path

EXCEL_SUFFIXES = (".xlsx", ".xlsm", ".xls")


class ExcelError(ValueError):
    """El archivo no se pudo leer como Excel o no tiene ninguna hoja con datos."""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
        return str(int(v))
    return str(v)


def _sheets_xlsx(path: Path) -> Iterator[tuple[str, Iterable[tuple]]]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - dependencia declarada en pyproject
        raise ExcelError("Para leer .xlsx falta la librería openpyxl (pip install openpyxl)") from exc
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - archivo dañado o que no es Excel: se informa al analista
        raise ExcelError(f"No se pudo abrir como Excel: {type(exc).__name__}: {exc}") from exc
    try:
        for ws in wb.worksheets:
            yield ws.title, ws.iter_rows(values_only=True)
    finally:
        wb.close()


def _sheets_xls(path: Path) -> Iterator[tuple[str, Iterable[tuple]]]:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise ExcelError("Para leer .xls (Excel antiguo) falta la librería xlrd (pip install xlrd)") from exc
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelError(f"No se pudo abrir como Excel antiguo (.xls): {type(exc).__name__}: {exc}") from exc

    def rows(sh) -> Iterator[tuple]:
        for r in range(sh.nrows):
            out = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    out.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    out.append(bool(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    out.append(None)
                else:
                    out.append(cell.value)
            yield tuple(out)

    try:
        for sh in book.sheets():
            yield sh.name, rows(sh)
    finally:
        book.release_resources()


def excel_to_csv(path: str | Path, target_for: Callable[[str], Path]) -> list[dict]:
    """Convierte cada hoja con datos en un CSV. `target_for(nombre_de_hoja)` devuelve la ruta donde escribirlo.

    Devuelve [{"sheet", "path", "rows", "columns"}] en el orden de las hojas. Lanza ExcelError si no hay ninguna hoja con datos.
    """
    path = Path(path)
    reader = _sheets_xls if path.suffix.lower() == ".xls" else _sheets_xlsx
    out = []
    for sheet, rows in reader(path):
        it = iter(rows)
        header = next((row for row in it if any(v not in (None, "") for v in row)), None)
        if header is None:
            continue  # hoja vacía
        width = max(i + 1 for i, v in enumerate(header) if v not in (None, ""))
        names, seen = [], {}
        for i, v in enumerate(header[:width]):
            name = _cell(v).strip() or f"col_{i + 1}"
            seen[name] = seen.get(name, 0) + 1
            names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
        target = target_for(sheet)
        n = 0
        with target.open("w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(names)
            for row in it:
                values = [_cell(v) for v in (list(row) + [None] * width)[:width]]
                if any(values):
                    w.writerow(values)
                    n += 1
        if n == 0:
            target.unlink()
            continue
        out.append({"sheet": sheet, "path": target, "rows": n, "columns": len(names)})
    if not out:
        raise ExcelError("El Excel no tiene ninguna hoja con datos (una cabecera y al menos una fila)")
    return out
