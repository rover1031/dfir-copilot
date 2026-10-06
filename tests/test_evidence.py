"""Entrega A: análisis con carpeta de evidencia propia. Lo esencial: lo que entra queda en una cadena de custodia con su SHA-256 y su origen;
ninguna ruta sale de la raíz de datos; un Excel se analiza por hojas con fechas legibles; un PDF se guarda pero no se analiza como log; y
confirmar la zona horaria no cambia ningún dato, solo deja constancia (y el informe dice con qué base)."""
import datetime
import json

import openpyxl
import pytest

from dfir_copilot.pipeline import Deps, run_pipeline
from dfir_copilot.projects import Project, ProjectError, ProjectSettings, browse, safe_name
from dfir_copilot.reporting import build_report
from dfir_copilot.synthetic import make_idor_dataset, write_csv
from dfir_copilot.timezone_status import TimezoneError, confirm_timezone, timezone_state


@pytest.fixture()
def world(tmp_path, monkeypatch):
    monkeypatch.setenv("DFIR_DATA_ROOT", str(tmp_path))
    monkeypatch.setenv("DFIR_PROJECTS_ROOT", str(tmp_path / "projects"))
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    return tmp_path, inbox


def new(tmp, **kw):
    return Project.create("Análisis Uno", root=tmp / "projects", base=tmp, ticket="INC-1234", **kw)


def xlsx(path, sheets: dict):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


# --- creación y nombres --------------------------------------------------------------------------------------------------

def test_un_analisis_nuevo_tiene_su_propia_carpeta_de_evidencia(world):
    tmp, _ = world
    p = new(tmp, description="firewall de borde")
    assert p.mode == "evidence" and p.ticket == "INC-1234" and p.description == "firewall de borde"
    assert p.source_dir == p.dir / "evidencia" and p.source_dir.is_dir() and p.files() == []


def test_los_proyectos_por_carpeta_siguen_igual(world):
    tmp, inbox = world
    p = Project.create("Viejo", inbox, root=tmp / "projects", base=tmp)
    assert p.mode == "folder" and p.source_dir == inbox.resolve()
    with pytest.raises(ProjectError, match="Nuevo análisis"):
        p.add_upload("a.csv", [b"x"])


@pytest.mark.parametrize("raw, safe", [("C:\\Users\\x\\..\\fw log.csv", "fw_log.csv"), ("../../etc/passwd.csv", "passwd.csv"),
                                       (".oculto.csv", "oculto.csv"), ("Tráfico Ñandú.xlsx", "Trafico_Nandu.xlsx"), ("", "archivo")])
def test_el_nombre_de_un_archivo_subido_nunca_lleva_rutas(raw, safe):
    assert safe_name(raw) == safe


# --- subir y traer del servidor ------------------------------------------------------------------------------------------

def test_subir_registra_hash_origen_y_quien_y_no_pisa_nombres(world):
    tmp, _ = world
    p = new(tmp)
    a = p.add_upload("fw.csv", [b"timestamp,src_ip\n", b"2026-10-04 10:00:00,10.0.0.1\n"], analyst="eder")
    b = p.add_upload("fw.csv", [b"otra cosa\n"], analyst="eder")
    assert [r["file"] for r in a + b] == ["fw.csv", "fw-2.csv"]
    rec = a[0]
    assert rec["action"] == "added" and rec["origin"] == "upload" and rec["analyst"] == "eder"
    assert rec["bytes"] == len(b"timestamp,src_ip\n2026-10-04 10:00:00,10.0.0.1\n")
    import hashlib
    assert rec["sha256"] == hashlib.sha256(b"timestamp,src_ip\n2026-10-04 10:00:00,10.0.0.1\n").hexdigest()


def test_un_tipo_no_admitido_o_demasiado_grande_no_deja_rastro(world):
    tmp, _ = world
    p = new(tmp)
    with pytest.raises(ProjectError, match="no admitido"):
        p.add_upload("malware.exe", [b"MZ"])
    with pytest.raises(ProjectError, match="límite"):
        p.add_upload("grande.csv", [b"x" * 600_000, b"x" * 600_000], max_bytes=1 << 20)
    assert list(p.source_dir.iterdir()) == [] and p.custody() == []


def test_traer_del_servidor_copia_o_enlaza_y_no_sale_de_la_raiz(world):
    tmp, inbox = world
    (inbox / "a.csv").write_text("x,y\n1,2\n")
    (inbox / "b.csv").write_text("x,y\n3,4\n")
    p = new(tmp)
    copy = p.add_from_server(inbox / "a.csv")[0]
    link = p.add_from_server(inbox / "b.csv", mode="link")[0]
    assert copy["mode"] == "copy" and not (p.source_dir / "a.csv").is_symlink()
    assert link["mode"] == "link" and (p.source_dir / "b.csv").is_symlink() and link["origin_path"] == str((inbox / "b.csv").resolve())
    outside = tmp.parent / "fuera.csv"
    outside.write_text("x\n")
    (inbox / "escape.csv").symlink_to(outside)                               # un enlace que apunta fuera de la raíz
    for bad in ("/etc/passwd", outside, inbox / "escape.csv", inbox / ".." / ".." / "fuera.csv"):
        with pytest.raises(ProjectError):
            p.add_from_server(bad)
    with pytest.raises(ProjectError, match="ya pertenece"):
        p.add_from_server(p.source_dir / "a.csv")


# --- Excel y PDF ---------------------------------------------------------------------------------------------------------

def test_un_excel_se_analiza_por_hojas_con_fechas_legibles(world):
    tmp, inbox = world
    xlsx(inbox / "reporte.xlsx", {
        "Tráfico": [[None], ["timestamp", "src_ip", "dst_port", "accion"],
                    [datetime.datetime(2026, 10, 4, 23, 2, 59), "10.0.0.1", 80.0, "ALLOW"], [None], [datetime.datetime(2026, 10, 4, 23, 3), "10.0.0.2", 443, "BLOCK"]],
        "Vacía": [],
        "Resumen": [["total"], [2]]})
    p = new(tmp)
    recs = p.add_from_server(inbox / "reporte.xlsx", analyst="eder")
    assert [r["action"] for r in recs] == ["added", "derived", "derived"]
    derived = recs[1]
    assert derived["sheet"] == "Tráfico" and derived["rows"] == 2 and derived["derived_from"]["sha256"] == recs[0]["sha256"]
    text = (p.source_dir / derived["file"]).read_text()
    assert text.splitlines() == ["timestamp,src_ip,dst_port,accion", "2026-10-04 23:02:59,10.0.0.1,80,ALLOW", "2026-10-04 23:03:00,10.0.0.2,443,BLOCK"]
    kinds = {f.name: (f.kind, f.supported) for f in p.files()}
    assert kinds["reporte.xlsx"] == ("excel", False) and kinds[derived["file"]] == ("log", True)


def test_un_xls_antiguo_tambien(world):
    xlwt = pytest.importorskip("xlwt", reason="xlwt solo hace falta para FABRICAR un .xls de prueba")
    tmp, inbox = world
    wb = xlwt.Workbook()
    sh = wb.add_sheet("datos")
    fmt = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
    for c, v in enumerate(["timestamp", "ip"]):
        sh.write(0, c, v)
    sh.write(1, 0, datetime.datetime(2026, 10, 4, 8, 0, 0), fmt)
    sh.write(1, 1, "10.0.0.9")
    wb.save(str(inbox / "viejo.xls"))
    p = new(tmp)
    recs = p.add_from_server(inbox / "viejo.xls")
    assert recs[1]["rows"] == 1
    assert (p.source_dir / recs[1]["file"]).read_text().splitlines() == ["timestamp,ip", "2026-10-04 08:00:00,10.0.0.9"]


def test_un_excel_danado_queda_como_evidencia_y_el_fallo_registrado(world):
    tmp, inbox = world
    (inbox / "roto.xlsx").write_bytes(b"esto no es un excel")
    recs = new(tmp).add_from_server(inbox / "roto.xlsx")
    assert [r["action"] for r in recs] == ["added", "derivation_failed"] and "Excel" in recs[1]["error"]


def test_un_pdf_se_guarda_como_documento_y_no_se_analiza(world):
    tmp, _ = world
    p = new(tmp)
    p.add_upload("informe CTI.pdf", [b"%PDF-1.4 contenido"])
    (f,) = p.files()
    assert f.kind == "document" and not f.supported and "no es un log" in f.reason


# --- custodia ------------------------------------------------------------------------------------------------------------

def test_la_custodia_detecta_ediciones_y_archivos_cambiados_o_borrados(world):
    tmp, _ = world
    p = new(tmp)
    p.add_upload("a.csv", [b"x\n1\n"])
    p.add_upload("b.csv", [b"x\n2\n"])
    assert p.verify_custody() == {"ok": True, "records": 2, "files": 2, "problems": []}
    (p.source_dir / "a.csv").write_text("x\n9\n")
    (p.source_dir / "b.csv").unlink()
    problems = p.verify_custody()["problems"]
    assert any("a.csv" in x and "cambió" in x for x in problems) and any("b.csv" in x and "ya no está" in x for x in problems)
    lines = p.custody_path.read_text().splitlines()
    rec = json.loads(lines[0])
    rec["analyst"] = "otro"
    p.custody_path.write_text("\n".join([json.dumps(rec)] + lines[1:]) + "\n")
    assert any("cadena" in x for x in p.verify_custody()["problems"])


# --- navegar el servidor -------------------------------------------------------------------------------------------------

def test_navegar_no_sale_de_la_raiz_ni_muestra_los_proyectos(world):
    tmp, inbox = world
    (inbox / "a.csv").write_text("x\n")
    (inbox / "x.exe").write_text("x\n")
    (inbox / ".oculto").write_text("x\n")
    new(tmp)
    root = browse("", base=tmp)
    assert [d["name"] for d in root["dirs"]] == ["inbox"] and root["parent"] is None        # "projects" no aparece
    sub = browse("inbox", base=tmp)
    assert {f["name"]: f["kind"] for f in sub["files"]} == {"a.csv": "log", "x.exe": "other"} and sub["parent"] == ""
    for bad in ("..", "../..", "/etc", "inbox/../../"):
        with pytest.raises(ProjectError):
            browse(bad, base=tmp)


# --- zona horaria --------------------------------------------------------------------------------------------------------

@pytest.fixture()
def analyzed(world):
    tmp, inbox = world
    p = Project.create("Zona", root=tmp / "projects", base=tmp,
                       settings=ProjectSettings(analyst="eder", timezone="America/Santiago", use_llm=False))
    rows, _ = make_idor_dataset()
    p.add_from_server(write_csv(inbox / "three_months.csv", rows), analyst="eder")
    st = run_pipeline(p, p.files()[0], Deps())
    assert st["steps"]["ingest"]["status"] in ("done", "needs_attention"), st
    return p.workspace(p.files()[0].case_id)


def test_confirmar_la_zona_deja_constancia_sin_tocar_los_datos(analyzed):
    ws = analyzed
    engine = ws.engine()
    before = ws.engine().manifest["output"]["sha256"]
    ledger = ws.ledger(engine)
    assert timezone_state(ledger, engine.manifest)["state"] == "unverified"
    with pytest.raises(TimezoneError, match="reingestar"):
        confirm_timezone(ledger, engine.manifest, "America/Bogota", "analyst_decision", "eder")
    with pytest.raises(TimezoneError, match="Base desconocida"):
        confirm_timezone(ledger, engine.manifest, "America/Santiago", "porque_si", "eder")
    confirm_timezone(ledger, engine.manifest, "America/Santiago", "analyst_decision", "eder", "acordado en el cierre")
    state = timezone_state(ws.ledger(ws.engine()), engine.manifest)
    assert state["state"] == "analyst_decision" and state["confirmation"]["confirmed_by"] == "eder"
    assert ws.engine().manifest["output"]["sha256"] == before and ws.verify().ok


def test_el_informe_dice_la_base_de_la_zona_tal_cual(analyzed):
    ws = analyzed
    md = build_report(ws, variant="compartible", lang="es", run_replay=False).markdown
    assert "SIN verificar" in md and "NO está verificada" in md
    engine = ws.engine()
    confirm_timezone(ws.ledger(engine), engine.manifest, "America/Santiago", "analyst_decision", "eder")
    md = build_report(ws, variant="compartible", lang="es", run_replay=False).markdown
    assert "decisión del analista, SIN confirmación externa (eder," in md and "sin confirmación externa. Si fuera otra" in md
    assert "NO está verificada" not in md
    en = build_report(ws, variant="compartible", lang="en", run_replay=False).markdown
    assert "WITHOUT external confirmation" in en
